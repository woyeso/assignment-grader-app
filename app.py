import os
import re
import logging
import json
import streamlit as st
import pdfplumber
from docx import Document
from transformers import AutoModelForCausalLM, AutoTokenizer
import torch

# Set up logging
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

# Paths to rubric files
P1_RUBRICS_PATH = os.path.join("data", "rubrics", "p1_rubrics.json")
P2_RUBRICS_PATH = os.path.join("data", "rubrics", "p2_rubrics.json")

# Load rubrics from JSON files
def load_rubrics(project_type):
    rubric_file = P1_RUBRICS_PATH if project_type.lower() == "group" else P2_RUBRICS_PATH
    try:
        with open(rubric_file, "r") as f:
            return json.load(f)
    except FileNotFoundError:
        raise FileNotFoundError(f"Rubric file not found: {rubric_file}")
    except json.JSONDecodeError:
        raise ValueError(f"Error decoding JSON from {rubric_file}")

# Load model and tokenizer
@st.cache(allow_output_mutation=True)
def load_model():
    model_name = "distilgpt2"  # Lightweight model
    hf_token = os.getenv("HF_TOKEN")

    try:
        # Try loading with fast tokenizer first
        tokenizer = AutoTokenizer.from_pretrained(
            model_name,
            token=hf_token if hf_token else None,
            use_fast=True
        )
    except Exception as e:
        logger.error(f"Failed to load fast tokenizer: {e}. Falling back to slow tokenizer.")
        # Fall back to slow tokenizer
        tokenizer = AutoTokenizer.from_pretrained(
            model_name,
            token=hf_token if hf_token else None,
            use_fast=False
        )

    # Set padding token for distilgpt2 (which doesn't have one by default)
    if tokenizer.pad_token is None:
        tokenizer.pad_token = tokenizer.eos_token
        logger.debug(f"Set pad_token to eos_token: {tokenizer.pad_token}")

    model = AutoModelForCausalLM.from_pretrained(
        model_name,
        torch_dtype=torch.float16,
        device_map="auto",
        token=hf_token if hf_token else None
    )

    # Ensure the model knows about the padding token
    model.config.pad_token_id = tokenizer.pad_token_id

    return model, tokenizer

model, tokenizer = load_model()

# Subcomponent mappings (same as original)
P1_SUBCOMPONENTS = {
    '1.1': 'Information of the Service Recipients Found:',
    '1.2': 'Information Related to the Use of AI in Teaching and Learning:',
    '1.3': 'Service Project Title and Topics:',
    '1.4': 'Specific Project Objectives:',
    '2.1': 'Design of AI-Related Ice-breaking Games:',
    '2.2': 'Tasks of Each Team Member:',
    '3.1': 'Specific STEM Elements Covered:',
    '3.2': 'Student Abilities to Strengthen:',
    '3.3': 'Potential Learning Hurdles of Students:',
    '3.4': 'Facilitating STEM and Overcoming Hurdles:',
    '4.1': 'List of Materials and Parts:',
    '4.2': 'List of Tools:'
}

P2_SUBCOMPONENTS = {
    '1.1': 'Specific Learning Objectives:',
    '1.2': 'Content of Each Teaching Kit:',
    '2.1': 'Describe the Design of Each Teaching Kit:',
    '2.2': 'How to Prepare (or Make) Each Item of Your Teaching Kit:',
    '2.3': 'Explain Why Students Will Learn and Play Happily:',
    '3.1': 'Draw a Diagram to Illustrate Task Breakdown:',
    '4.1': 'How to Introduce the Specific Topic(s) to Arouse Interest in STEM:',
    '4.2': 'How to Identify and Overcome Learning Hurdles:',
    '5.1': 'How to React to Potential Uncertainties:',
    '5.2': 'How to Self-Evaluate Performance and Make Improvements:'
}

# Text extraction functions (unchanged)
def extract_text_between_strings(text, start_keyword, end_keyword):
    try:
        extracted_text = ""
        start_match = re.search(start_keyword, text, re.MULTILINE)
        if not start_match:
            logger.debug(f"Start keyword '{start_keyword}' not found.")
            return "Not Found"
        
        start_index = start_match.end()
        end_match = re.search(end_keyword, text, re.MULTILINE)
        if end_match and end_match.start() > start_match.start():
            end_index = end_match.start()
            extracted_text = text[start_index:end_index].strip()
        else:
            extracted_text = text[start_index:].strip()
        
        if not extracted_text:
            logger.debug(f"End keyword '{end_keyword}' not found or no content extracted.")
            return "Not Found"

        lines = extracted_text.split('\n')
        formatted_lines = []
        bullet_pattern = re.compile(r'^\s*(\d+\.|\•|-|◦|➢)\s*(.+)$')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            bullet_match = bullet_pattern.match(line)
            if bullet_match:
                bullet, text = bullet_match.groups()
                formatted_lines.append(f"{bullet} {text}")
            else:
                formatted_lines.append(line)
        cleaned_text = "\n".join(formatted_lines).strip()
        cleaned_text = re.sub(r'\s+', ' ', cleaned_text.replace('\n', '\n '))
        return cleaned_text.replace("XYZ students", "Hong Chi students")

    except Exception as e:
        logger.error(f"Error extracting text: {e}")
        return f"Error: {e}"

def extract_text_from_pdf(filepath, assignment_type='P1'):
    results = {}
    subcomponents = P1_SUBCOMPONENTS if assignment_type == 'P1' else P2_SUBCOMPONENTS
    sorted_codes = sorted(subcomponents.keys(), key=lambda x: [int(n) for n in x.split('.')])
    
    with pdfplumber.open(filepath) as pdf:
        text = ""
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"
    
    for i, code in enumerate(sorted_codes):
        start_keyword = r"^{}\s*[.:]?\s*".format(re.escape(code))
        if i + 1 < len(sorted_codes):
            end_keyword = r"^{}\s*[.:]?\s*".format(re.escape(sorted_codes[i + 1]))
        else:
            end_keyword = r"^5\.\s*" if assignment_type == 'P1' else r"^6\.\s*"
        
        logger.debug(f"Extracting section {code} with start_keyword={start_keyword}, end_keyword={end_keyword}")
        content = extract_text_between_strings(text, start_keyword, end_keyword)
        results[code] = {
            "title": subcomponents[code],
            "content": content
        }

    return results

def extract_text_from_docx(filepath, assignment_type='P1'):
    try:
        doc = Document(filepath)
        elements = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                style = para.style.name
                elements.append(('paragraph', text, style))
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    table_text.append(" ".join(row_text))
            if table_text:
                elements.append(('table', "\n".join(table_text), 'Table'))
        
        logger.debug(f"Extracted {len(elements)} elements from DOCX")
        
        results = {}
        subcomponents = P1_SUBCOMPONENTS if assignment_type == 'P1' else P2_SUBCOMPONENTS
        sorted_codes = sorted(subcomponents.keys(), key=lambda x: [int(n) for n in x.split('.')])
        
        current_section = None
        section_content = []
        section_pattern = re.compile(r'^\s*(\d+\.\d+\.?)\s*[.:]?\s*(.*)?$')
        end_pattern = re.compile(r'^\s*5\.\s*' if assignment_type == 'P1' else r'^\s*6\.\s*')
        bullet_pattern = re.compile(r'^\s*(\d+\.|\•|-|◦|➢)\s*(.+)$')
        
        for i, (elem_type, text, style) in enumerate(elements):
            logger.debug(f"Processing element {i}: type={elem_type}, style={style}, text={text[:100]}...")
            
            lines = text.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                section_match = section_pattern.match(line)
                if section_match:
                    code, title = section_match.groups()
                    code = code.rstrip('.')
                    if current_section and current_section in subcomponents:
                        formatted_lines = []
                        for content_line in section_content:
                            bullet_match = bullet_pattern.match(content_line)
                            if bullet_match:
                                bullet, text = bullet_match.groups()
                                formatted_lines.append(f"{bullet} {text}")
                            else:
                                formatted_lines.append(content_line)
                        cleaned_content = "\n".join(formatted_lines).strip()
                        cleaned_content = re.sub(r'\s+', ' ', cleaned_content.replace('\n', '\n '))
                        cleaned_content = cleaned_content.replace("XYZ students", "Hong Chi students")
                        results[current_section] = {
                            "title": subcomponents[current_section],
                            "content": cleaned_content if cleaned_content else "Not Found"
                        }
                    current_section = code
                    section_content = []
                    if title:
                        section_content.append(title)
                    logger.debug(f"Started section {code} at element {i}")
                    continue
                
                end_match = end_pattern.match(line)
                if end_match and current_section:
                    formatted_lines = []
                    for content_line in section_content:
                        bullet_match = bullet_pattern.match(content_line)
                        if bullet_match:
                            bullet, text = bullet_match.groups()
                            formatted_lines.append(f"{bullet} {text}")
                        else:
                            formatted_lines.append(content_line)
                    cleaned_content = "\n".join(formatted_lines).strip()
                    cleaned_content = re.sub(r'\s+', ' ', cleaned_content.replace('\n', '\n '))
                    cleaned_content = cleaned_content.replace("XYZ students", "Hong Chi students")
                    results[current_section] = {
                        "title": subcomponents[current_section],
                        "content": cleaned_content if cleaned_content else "Not Found"
                    }
                    current_section = None
                    section_content = []
                    logger.debug(f"Ended section at element {i} with end marker")
                    continue
                
                if current_section:
                    if style.startswith('List') or bullet_pattern.match(line):
                        bullet_match = bullet_pattern.match(line)
                        if bullet_match:
                            bullet, text = bullet_match.groups()
                            section_content.append(f"{bullet} {text}")
                        else:
                            section_content.append(f"- {line}")
                    else:
                        section_content.append(line)
        
        if current_section and current_section in subcomponents:
            formatted_lines = []
            for content_line in section_content:
                bullet_match = bullet_pattern.match(content_line)
                if bullet_match:
                    bullet, text = bullet_match.groups()
                    formatted_lines.append(f"{bullet} {text}")
                else:
                    formatted_lines.append(content_line)
            cleaned_content = "\n".join(formatted_lines).strip()
            cleaned_content = re.sub(r'\s+', ' ', cleaned_content.replace('\n', '\n '))
            cleaned_content = cleaned_content.replace("XYZ students", "Hong Chi students")
            results[current_section] = {
                "title": subcomponents[current_section],
                "content": cleaned_content if cleaned_content else "Not Found"
            }
        
        for code in sorted_codes:
            if code not in results:
                results[code] = {
                    "title": subcomponents[code],
                    "content": "Not Found"
                }
                logger.debug(f"Subcomponent {code} not found in DOCX")
        
        return results
    
    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {e}")
        return {}

# Function to evaluate submission using the model
def evaluate_submission(subcomponent, project_type, rubric, submission, school_name):
    prompt = (
        f"Can you evaluate my project submission for Subcomponent {subcomponent} in a {project_type} project (P1 for group, P2 for individual).\n"
        f"Here is the rubric: {rubric}. Evaluate the submission against each rubric criterion. Focus on the rubric criteria as the primary basis for your evaluation.\n"
        f"My submission is {submission}.\n\n"
        f"If a school name is provided, use it in your evaluation: {school_name}. If no school name is provided, refer to the students generically as 'students'.\n"
        f"Do not use the placeholder 'XYZ students' in your evaluation, as it was used during training but should be replaced with the specific school name or 'students'.\n\n"
        f"Summarize the strengths of the submission (what it does well according to the rubric).\n"
        f"Summarize the weaknesses of the submission (where it falls short according to the rubric).\n"
        f"Provide specific suggestions for improvement to help the student improve their submission.\n\n"
        f"Give me an overall mark out of 10, and don't be too strict. Ensure you provide the score in the format: <Overall Mark: X/10>. Do not omit the score and follow format of X/10."
    )
    inputs = tokenizer(prompt, return_tensors="pt", padding=True, truncation=True, max_length=512)
    inputs = {k: v.to(model.device) for k, v in inputs.items()}
    
    with torch.no_grad():
        outputs = model.generate(
            input_ids=inputs["input_ids"],
            attention_mask=inputs["attention_mask"],
            max_new_tokens=256,
            temperature=0.7,
            top_p=0.9,
            do_sample=True,
            pad_token_id=tokenizer.pad_token_id  # Explicitly set pad_token_id for generation
        )
    feedback = tokenizer.decode(outputs[0], skip_special_tokens=True)
    return feedback

# Streamlit app
st.title("Assignment Grader App")

# File upload
uploaded_file = st.file_uploader("Upload PDF/DOCX", type=["pdf", "docx"])
project_type = st.selectbox("Project Type", ["Group (P1)", "Individual (P2)"])
school_name = st.text_input("School Name (Optional)")
group_number = st.text_input("Group Number (Optional)")

# Manual text input as fallback
manual_text = st.text_area("Or enter your submission text manually (optional)")

if st.button("Evaluate"):
    if uploaded_file or manual_text:
        # Load rubrics
        project_type_short = "Group" if project_type == "Group (P1)" else "Individual"
        project = "P1" if project_type == "Group (P1)" else "P2"
        try:
            rubrics = load_rubrics(project_type_short)
        except Exception as e:
            st.error(f"Error loading rubrics: {str(e)}")
            st.stop()

        # Extract text from file or use manual input
        submission_dict = {}
        if uploaded_file:
            with open("/tmp/uploaded_file", "wb") as f:
                f.write(uploaded_file.read())
            
            if uploaded_file.name.endswith(".pdf"):
                results = extract_text_from_pdf("/tmp/uploaded_file", project)
            else:
                results = extract_text_from_docx("/tmp/uploaded_file", project)
            os.remove("/tmp/uploaded_file")
            
            for subcomponent, data in results.items():
                if data["content"] != "Not Found":
                    submission_dict[subcomponent] = data["content"]
        else:
            submission_dict["1.1"] = manual_text  # Simplified for manual input; adjust as needed

        if not submission_dict:
            st.error("No text extracted from the file or provided manually.")
            st.stop()

        # Evaluate submissions
        evaluations = []
        total_score = 0
        total_weight = 0

        with st.spinner("Evaluating submission..."):
            for rubric in rubrics:
                subcomponent = rubric["subcomponent"]
                if subcomponent not in submission_dict:
                    continue

                submission = submission_dict[subcomponent]
                evaluation = evaluate_submission(
                    subcomponent,
                    project_type_short,
                    rubric["criteria"],
                    submission,
                    school_name if school_name else "Not provided"
                )

                if school_name:
                    evaluation = evaluation.replace("XYZ students", f"{school_name} students")
                else:
                    evaluation = evaluation.replace("XYZ students", "students")

                score_match = re.search(r"Overall Mark:\s*([\d.]+)(?:\s*/\s*10)?", evaluation, re.IGNORECASE)
                score = float(score_match.group(1)) if score_match else 0

                weight = rubric.get("weight", 1.0)
                total_score += score * weight
                total_weight += weight

                evaluations.append({
                    "subcomponent": subcomponent,
                    "evaluation": evaluation,
                    "score": score,
                    "weight": weight
                })

        # Calculate final grade
        final_grade = (total_score / total_weight) * 10 if total_weight > 0 else 0
        final_grade = round(final_grade, 2)

        # Display results
        group_display = f" {group_number}" if group_number else ""
        summary = f"**Summary of Evaluations for {project} Project (Group{group_display})**\n\n"
        separator = "********************************************************************\n"
        for i, eval in enumerate(evaluations):
            summary += f"**Subcomponent {eval['subcomponent']} (Weight: {eval['weight']*100}%)**\n"
            summary += eval["evaluation"]
            summary += "\n\n"
            if i < len(evaluations) - 1:
                summary += separator

        summary += f"**Final Total Grade: {final_grade}%**"

        st.subheader("Evaluation Results")
        st.markdown(summary)
    else:
        st.error("Please upload a file or enter text manually.")